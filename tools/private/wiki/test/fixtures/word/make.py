# Run: python3 -m venv v && v/bin/pip install python-docx && v/bin/python test/fixtures/word/make.py test/fixtures/word/Livret-accueil.docx
# Writes test/fixtures/word/Livret-accueil.docx with python-docx (MIT), whose
# default template is a Word document: styles, numbering and relations as
# Word writes them.
import sys, struct, zlib
from docx import Document
from docx.shared import Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def png(w, h):
    raw = b"".join(b"\x00" + b"\xff\x99\x33" * w for _ in range(h))
    def chunk(t, d): return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xffffffff)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")

out = sys.argv[1]
open(out + ".png", "wb").write(png(12, 8))
d = Document()
d.core_properties.title = "Livret d’accueil (propriétés)"
d.add_heading("Livret d’accueil", level=0)  # the Title style
p = d.add_paragraph("Bienvenue chez ")
p.add_run("Lumen & Co").bold = True
p.add_run(" ! Ce livret explique ")
p.add_run("comment nous travaillons").italic = True
p.add_run(".")
d.add_heading("Votre première semaine", level=1)
d.add_paragraph("Récupérer votre badge", style="List Bullet")
d.add_paragraph("Lire le règlement intérieur", style="List Bullet")
d.add_paragraph("Signer la charte informatique", style="List Bullet 2")
d.add_heading("Horaires", level=2)
d.add_paragraph("Arrivée avant 10 h", style="List Number")
d.add_paragraph("Départ après 16 h", style="List Number")
t = d.add_table(rows=3, cols=2)
for r, row in enumerate([("Pour", "Demander à"), ("Congés", "Votre responsable"), ("Wi-Fi", "Tom")]):
    for c, text in enumerate(row):
        t.cell(r, c).text = text
d.add_paragraph("Le plan des locaux :")
d.add_picture(out + ".png", width=Inches(1.2))
# A link to the web, as Word writes one.
p = d.add_paragraph("Textes officiels : ")
rid = d.part.relate_to("https://www.service-public.fr/", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
r = OxmlElement("w:r"); tt = OxmlElement("w:t"); tt.text = "service-public.fr"; r.append(tt); h.append(r); p._p.append(h)
p = d.add_paragraph("Une ligne")
p.add_run().add_break()
p.add_run("puis la suivante, ")
p.add_run("barrée").font.strike = True
d.add_paragraph("Citation du fondateur.", style="Quote")
d.save(out)
